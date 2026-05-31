import openpyxl, sys, datetime
sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook(r'docs\Reporte de ataques.xlsx', data_only=True)
ws = wb.active

animal_names = ['Rat','Spider','Snake','Bat','Boar','Wolf','Bear','Croc','Tiger','Elephant']
dangerous = {'Bear','Croc','Tiger','Elephant'}

rows = []
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    if any(v is not None for v in row):
        rows.append(row)

check_time = datetime.time(20, 20)
check_dt   = datetime.datetime.combine(datetime.date.today(), check_time)

print(f'Total oasis: {len(rows)}\n')
print(f'{"Gap":>8}  {"Total":>6}  Composicion')
print('-'*75)

for row in rows:
    last_t = row[0]
    if isinstance(last_t, datetime.time):
        last_dt = datetime.datetime.combine(datetime.date.today(), last_t)
    else:
        continue
    gap_min = int((check_dt - last_dt).total_seconds() / 60)
    animals  = row[2:]
    total    = sum(a for a in animals if a)
    comp     = ', '.join(f'{int(a)} {n}' for a, n in zip(animals, animal_names) if a)
    flag     = ''
    for a, n in zip(animals, animal_names):
        if a and n in dangerous and a >= 1:
            flag = f'  *** {n.upper()} ***'
            break
    print(f'{gap_min:>6}m  {total:>6.0f}  {comp}{flag}')

totals = []
for row in rows:
    if isinstance(row[0], datetime.time):
        totals.append(sum(a for a in row[2:] if a))

print(f'\n--- RESUMEN ---')
print(f'Minimo:  {min(totals):.0f} animales')
print(f'Maximo:  {max(totals):.0f} animales')
print(f'Media:   {sum(totals)/len(totals):.1f} animales')
print(f'Oasis con >20 animales: {sum(1 for t in totals if t > 20)}')
print(f'Oasis con >40 animales: {sum(1 for t in totals if t > 40)}')
dangerous_count = 0
for row in rows:
    for a, n in zip(row[2:], animal_names):
        if a and n in dangerous:
            dangerous_count += 1
            break
print(f'Oasis con animales peligrosos (oso/croc/tigre/elefante): {dangerous_count}')
