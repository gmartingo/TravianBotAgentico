import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook(r'docs\Copia de Copy of Oasis farming from a nerd.xlsx', data_only=True)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'\n{"="*60}')
    print(f'HOJA: {sheet_name}  ({ws.max_row}x{ws.max_column})')
    print('='*60)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        if any(v is not None for v in row):
            print(row)
